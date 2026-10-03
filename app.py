import os
import io
import json
import re
import uuid
from datetime import datetime
from flask import Flask, request, jsonify, render_template_string
import gspread
from google.oauth2.service_account import Credentials
import boto3
from botocore.client import Config
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

app = Flask(__name__)

# --- Конфигурация ---
BUCKET_NAME = os.environ['BUCKET_NAME']
GOOGLE_SHEET_ID = os.environ['GOOGLE_SHEET_ID']
YC_KEY_ID = os.environ['YC_KEY_ID']
YC_PRIVATE_KEY = os.environ['YC_PRIVATE_KEY']

# Google Sheets auth
google_creds = json.loads(os.environ['GOOGLE_SERVICE_ACCOUNT_JSON'])
gc = gspread.service_account_from_dict(google_creds)
sh = gc.open_by_key(GOOGLE_SHEET_ID)
worksheet_objects = sh.worksheet('объекты')
worksheet_reestr = sh.worksheet('реестр')
worksheet_owners = sh.worksheet('справочник_собственников')
worksheet_template = sh.worksheet('шаблон_кабельного_журнала')

# S3/Yandex Object Storage
session = boto3.session.Session()
s3 = session.client(
    's3',
    endpoint_url='https://storage.yandexcloud.net',
    aws_access_key_id=YC_KEY_ID,
    aws_secret_access_key=YC_PRIVATE_KEY,
    config=Config(signature_version='s3v4')
)

# --- HTML шаблон ---
HTML_TEMPLATE = '''
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Фотофиксация кабелей</title>
    <style>
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body { font-family: Arial, sans-serif; background: #f0f2f5; padding: 20px; }
        .container { max-width: 800px; margin: 0 auto; }
        .card { background: white; border-radius: 12px; padding: 24px; margin: 16px 0; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }
        h1 { color: #1a73e8; margin-bottom: 24px; }
        h2 { color: #333; margin-bottom: 16px; }
        input, select, textarea { width: 100%; padding: 10px; margin: 8px 0; border: 1px solid #ddd; border-radius: 6px; font-size: 14px; }
        .birka-input { font-size: 32px; height: 80px; letter-spacing: 10px; font-family: 'Courier New', monospace; text-align: center; }
        .photo-upload { border: 2px dashed #1a73e8; padding: 20px; text-align: center; margin: 8px 0; cursor: pointer; }
        .photo-upload:hover { background: #e8f0fe; }
        .btn { background: #1a73e8; color: white; padding: 12px 24px; border: none; border-radius: 6px; cursor: pointer; font-size: 16px; margin: 8px 4px; }
        .btn:hover { opacity: 0.9; }
        .btn-danger { background: #dc3545; }
        .btn-warning { background: #ffc107; color: #333; }
        .btn-success { background: #28a745; }
        .error { background: #f8d7da; padding: 12px; border-radius: 6px; margin: 8px 0; display: none; }
        .success { background: #d4edda; padding: 12px; border-radius: 6px; margin: 8px 0; display: none; }
        .info { background: #cce5ff; padding: 12px; border-radius: 6px; margin: 8px 0; }
        .object-info { background: #f8f9fa; padding: 12px; border-radius: 6px; margin: 8px 0; }
        .hidden { display: none; }
        .checkbox-group { margin: 12px 0; }
        .checkbox-group label { margin-right: 20px; }
        .autocomplete { position: relative; }
        .autocomplete-items { position: absolute; background: white; border: 1px solid #ddd; width: 100%; max-height: 200px; overflow-y: auto; z-index: 1000; }
        .autocomplete-items div { padding: 10px; cursor: pointer; }
        .autocomplete-items div:hover { background: #e8f0fe; }
    </style>
</head>
<body>
    <div class="container">
        <h1>📸 Фотофиксация кабелей</h1>
        
        <!-- Шаг 1: Поиск объекта -->
        <div class="card">
            <h2>Шаг 1: Выберите объект</h2>
            <input type="text" id="objectSearch" placeholder="Введите номер объекта" oninput="searchObject(this.value)">
            <div id="objectInfo" class="hidden"></div>
        </div>
        
        <!-- Шаг 2: Типы фото -->
        <div class="card hidden" id="photoTypes">
            <h2>Шаг 2: Типы фотографий</h2>
            <p>Сделайте следующие фотографии:</p>
            <div class="checkbox-group">
                <label><input type="checkbox" class="photo-type" value="1"> 1. Адрес дома (необязательно)</label><br>
                <label><input type="checkbox" class="photo-type" value="2"> 2. Общий вид крыши</label><br>
                <label><input type="checkbox" class="photo-type" value="3"> 3. Зона размещения кабелей</label><br>
                <label><input type="checkbox" class="photo-type" value="4"> 4. Кабель с биркой</label>
            </div>
            <button class="btn" onclick="showCableForm()">Далее →</button>
        </div>
        
        <!-- Шаг 3: Форма кабеля -->
        <div class="card hidden" id="cableForm">
            <h2>Шаг 3: Данные кабеля</h2>
            
            <label>Номер бирки (6 цифр):</label>
            <input type="text" id="birkaNumber" class="birka-input" maxlength="6" placeholder="000000" oninput="checkBirka(this.value)">
            <div id="birkaError" class="error"></div>
            <div id="birkaConfirm" class="hidden">
                <button class="btn btn-success" onclick="confirmBirka()">✅ Подтвердить, номер верный</button>
                <button class="btn btn-warning" onclick="editBirka()">✏️ Исправить номер</button>
            </div>
            
            <div id="cableDetails" class="hidden">
                <h3>Точка Б</h3>
                <label>Точка Б1 (обязательно):</label>
                <input type="text" id="pointB1" placeholder="Адрес точки Б1">
                
                <label>Тип точки Б:</label>
                <select id="pointBType">
                    <option value="МКД">МКД</option>
                    <option value="Административное здание">Административное здание</option>
                    <option value="Прочее">Прочее</option>
                    <option value="Опора">Опора</option>
                </select>
                
                <div class="checkbox-group">
                    <label><input type="checkbox" id="isTransit" onchange="toggleB2()"> Транзитный кабель</label>
                </div>
                
                <div id="b2Field" class="hidden">
                    <label>Точка Б2 (для транзита):</label>
                    <input type="text" id="pointB2" placeholder="Адрес точки Б2">
                </div>
                
                <label>Собственник:</label>
                <div class="autocomplete">
                    <input type="text" id="owner" placeholder="Начните вводить или введите цифру" oninput="autocompleteOwner(this.value)">
                    <div id="ownerSuggestions" class="autocomplete-items"></div>
                </div>
                
                <div id="commentSection">
                    <label>Комментарий:</label>
                    <textarea id="comment" rows="3" placeholder="Дополнительный комментарий"></textarea>
                </div>
                
                <div class="photo-upload" onclick="document.getElementById('cablePhoto').click()">
                    📷 Загрузить фото кабеля с биркой (макс 30 шт)
                </div>
                <input type="file" id="cablePhoto" multiple accept="image/*" style="display:none" onchange="previewPhotos(event)">
                <div id="photoPreview"></div>
                
                <button class="btn" onclick="saveCable()">💾 Сохранить кабель</button>
            </div>
        </div>
        
        <!-- Кнопка "Кабели не выявлены" -->
        <div class="card hidden" id="noCableSection">
            <button class="btn btn-danger" onclick="noCableFound()">🚫 Кабели не выявлены</button>
        </div>
        
        <!-- Результат -->
        <div id="result" class="hidden"></div>
    </div>
    
    <script>
        let currentObject = null;
        let currentBirka = '';
        let birkaConfirmed = false;
        let photos = [];
        
        function searchObject(val) {
            if (val.length < 1) {
                document.getElementById('objectInfo').classList.add('hidden');
                return;
            }
            
            fetch('/search_object', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({query: val})
            })
            .then(r => r.json())
            .then(data => {
                if (data.found) {
                    currentObject = data.object;
                    const info = document.getElementById('objectInfo');
                    info.innerHTML = `
                        <div class="object-info">
                            <strong>Объект найден!</strong><br>
                            Адрес: ${data.object.address}<br>
                            Ответственный: ${data.object.responsible}<br>
                            Тип здания: ${data.object.building_type}
                        </div>
                    `;
                    info.classList.remove('hidden');
                    document.getElementById('photoTypes').classList.remove('hidden');
                    document.getElementById('noCableSection').classList.remove('hidden');
                } else {
                    document.getElementById('objectInfo').innerHTML = '<div class="error" style="display:block">Объект не найден</div>';
                    document.getElementById('objectInfo').classList.remove('hidden');
                    document.getElementById('photoTypes').classList.add('hidden');
                    document.getElementById('noCableSection').classList.add('hidden');
                }
            });
        }
        
        function showCableForm() {
            document.getElementById('cableForm').classList.remove('hidden');
        }
        
        function checkBirka(val) {
            if (val.length === 6) {
                currentBirka = val;
                fetch('/check_birka', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({birka: val})
                })
                .then(r => r.json())
                .then(data => {
                    const errorDiv = document.getElementById('birkaError');
                    const confirmDiv = document.getElementById('birkaConfirm');
                    if (data.duplicate) {
                        errorDiv.innerHTML = `⚠️ Бирка ${val} уже существует!<br>Объект: ${data.object_number}<br>Дата: ${data.date}`;
                        errorDiv.style.display = 'block';
                        confirmDiv.classList.remove('hidden');
                        birkaConfirmed = false;
                    } else {
                        errorDiv.style.display = 'none';
                        confirmDiv.classList.add('hidden');
                        birkaConfirmed = true;
                        document.getElementById('cableDetails').classList.remove('hidden');
                    }
                });
            }
        }
        
        function confirmBirka() {
            birkaConfirmed = true;
            document.getElementById('birkaError').style.display = 'none';
            document.getElementById('birkaConfirm').classList.add('hidden');
            document.getElementById('cableDetails').classList.remove('hidden');
        }
        
        function editBirka() {
            birkaConfirmed = false;
            document.getElementById('birkaError').style.display = 'none';
            document.getElementById('birkaConfirm').classList.add('hidden');
            document.getElementById('cableDetails').classList.add('hidden');
            document.getElementById('birkaNumber').value = '';
            document.getElementById('birkaNumber').focus();
        }
        
        function toggleB2() {
            if (document.getElementById('isTransit').checked) {
                document.getElementById('b2Field').classList.remove('hidden');
                document.getElementById('comment').value = 'Транзит ' + document.getElementById('pointB1').value + ' - ' + document.getElementById('pointB2').value;
            } else {
                document.getElementById('b2Field').classList.add('hidden');
                document.getElementById('comment').value = '';
            }
        }
        
        function autocompleteOwner(val) {
            if (val.length < 1) {
                document.getElementById('ownerSuggestions').innerHTML = '';
                return;
            }
            
            fetch('/search_owner', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({query: val})
            })
            .then(r => r.json())
            .then(data => {
                let html = '';
                data.owners.forEach(o => {
                    html += `<div onclick="selectOwner('${o.name}')">${o.name}</div>`;
                });
                document.getElementById('ownerSuggestions').innerHTML = html;
            });
        }
        
        function selectOwner(name) {
            document.getElementById('owner').value = name;
            document.getElementById('ownerSuggestions').innerHTML = '';
        }
        
        function previewPhotos(event) {
            photos = Array.from(event.target.files);
            const preview = document.getElementById('photoPreview');
            preview.innerHTML = `<div class="info">Выбрано ${photos.length} фото</div>`;
        }
        
        function saveCable() {
            if (!birkaConfirmed) {
                alert('Пожалуйста, подтвердите номер бирки');
                return;
            }
            
            if (photos.length === 0) {
                alert('Загрузите хотя бы одно фото');
                return;
            }
            
            if (photos.length > 30) {
                alert('Максимум 30 фото за раз');
                return;
            }
            
            const formData = new FormData();
            formData.append('object_number', currentObject.number);
            formData.append('object_address', currentObject.address);
            formData.append('object_type', currentObject.building_type);
            formData.append('object_lat', currentObject.lat);
            formData.append('object_lon', currentObject.lon);
            formData.append('birka', currentBirka);
            formData.append('point_b1', document.getElementById('pointB1').value);
            formData.append('point_b2', document.getElementById('pointB2').value);
            formData.append('point_b_type', document.getElementById('pointBType').value);
            formData.append('owner', document.getElementById('owner').value);
            formData.append('is_transit', document.getElementById('isTransit').checked ? 'ДА' : '');
            formData.append('comment', document.getElementById('comment').value);
            
            photos.forEach(photo => formData.append('photos', photo));
            
            fetch('/upload', {
                method: 'POST',
                body: formData
            })
            .then(r => r.json())
            .then(data => {
                if (data.success) {
                    document.getElementById('result').innerHTML = '<div class="success">✅ Данные сохранены!</div>';
                    document.getElementById('result').classList.remove('hidden');
                    resetForm();
                } else {
                    document.getElementById('result').innerHTML = `<div class="error">❌ Ошибка: ${data.error}</div>`;
                    document.getElementById('result').classList.remove('hidden');
                }
            });
        }
        
        function noCableFound() {
            if (!currentObject) return;
            
            fetch('/no_cable', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({
                    object_number: currentObject.number,
                    object_address: currentObject.address,
                    object_type: currentObject.building_type,
                    object_lat: currentObject.lat,
                    object_lon: currentObject.lon
                })
            })
            .then(r => r.json())
            .then(data => {
                if (data.success) {
                    document.getElementById('result').innerHTML = '<div class="success">✅ Данные сохранены (кабели не выявлены)</div>';
                    document.getElementById('result').classList.remove('hidden');
                }
            });
        }
        
        function resetForm() {
            currentObject = null;
            currentBirka = '';
            birkaConfirmed = false;
            photos = [];
            document.getElementById('objectSearch').value = '';
            document.getElementById('objectInfo').classList.add('hidden');
            document.getElementById('photoTypes').classList.add('hidden');
            document.getElementById('cableForm').classList.add('hidden');
            document.getElementById('noCableSection').classList.add('hidden');
            document.getElementById('birkaNumber').value = '';
            document.getElementById('pointB1').value = '';
            document.getElementById('pointB2').value = '';
            document.getElementById('owner').value = '';
            document.getElementById('comment').value = '';
            document.getElementById('isTransit').checked = false;
            document.getElementById('b2Field').classList.add('hidden');
            document.getElementById('photoPreview').innerHTML = '';
        }
    </script>
</body>
</html>
'''

# --- API endpoints ---
@app.route('/')
def index():
    return render_template_string(HTML_TEMPLATE)

@app.route('/search_object', methods=['POST'])
def search_object():
    data = request.get_json()
    query = data.get('query', '').strip()
    
    try:
        objects = worksheet_objects.get_all_records()
        for obj in objects:
            if str(obj.get('№ объекта', '')) == query:
                return jsonify({
                    'found': True,
                    'object': {
                        'number': str(obj['№ объекта']),
                        'address': obj['Адрес'],
                        'responsible': obj['Ответственный'],
                        'building_type': obj['Тип здания'],
                        'lat': str(obj.get('Широта', '')),
                        'lon': str(obj.get('Долгота', ''))
                    }
                })
        return jsonify({'found': False})
    except Exception as e:
        return jsonify({'found': False, 'error': str(e)})

@app.route('/check_birka', methods=['POST'])
def check_birka():
    data = request.get_json()
    birka = data.get('birka', '').strip()
    
    try:
        records = worksheet_reestr.get_all_values()
        for row in records[1:]:  # skip header
            if len(row) > 9 and row[9] == birka:  # column J is ID кабеля
                return jsonify({
                    'duplicate': True,
                    'object_number': row[0],
                    'date': row[1]
                })
        return jsonify({'duplicate': False})
    except Exception as e:
        return jsonify({'duplicate': False, 'error': str(e)})

@app.route('/search_owner', methods=['POST'])
def search_owner():
    data = request.get_json()
    query = data.get('query', '').strip()
    
    try:
        owners = worksheet_owners.get_all_values()
        results = []
        for row in owners:
            if len(row) >= 2:
                if query.isdigit() and row[0] == query:
                    results.append({'name': row[1]})
                elif query.lower() in row[1].lower():
                    results.append({'name': row[1]})
        return jsonify({'owners': results[:5]})
    except Exception as e:
        return jsonify({'owners': []})

@app.route('/upload', methods=['POST'])
def upload():
    try:
        object_number = request.form.get('object_number')
        object_address = request.form.get('object_address')
        object_type = request.form.get('object_type')
        object_lat = request.form.get('object_lat')
        object_lon = request.form.get('object_lon')
        birka = request.form.get('birka')
        point_b1 = request.form.get('point_b1')
        point_b2 = request.form.get('point_b2')
        point_b_type = request.form.get('point_b_type')
        owner = request.form.get('owner')
        is_transit = request.form.get('is_transit')
        comment = request.form.get('comment')
        
        photos = request.files.getlist('photos')
        
        # Save photos to Yandex Cloud
        for idx, photo in enumerate(photos):
            safe_address = secure_filename(object_address[:30])
            filename = f"{object_number}_{safe_address}_Бирка_{birka}_{idx+1}.jpg"
            s3.upload_fileobj(photo, BUCKET_NAME, filename)
        
        # Create/update cable journal
        journal_filename = f"{object_number}_{safe_address}_Кабельный журнал.xlsx"
        
        # Load template
        template_data = worksheet_template.get_all_values()
        wb = load_workbook()
        ws = wb.active
        
        # Copy template header (rows 1-6)
        for i in range(min(6, len(template_data))):
            for j in range(len(template_data[i])):
                ws.cell(row=i+1, column=j+1, value=template_data[i][j])
        
        # Add data starting from row 7
        row = 7
        ws.cell(row=row, column=1, value=1)  # №
        ws.cell(row=row, column=2, value=object_address)  # Точка А
        ws.cell(row=row, column=3, value=object_type)  # Тип А
        ws.cell(row=row, column=4, value=object_lat)  # Широта
        ws.cell(row=row, column=5, value=object_lon)  # Долгота
        ws.cell(row=row, column=6, value=point_b1)  # Точка Б
        ws.cell(row=row, column=7, value=point_b_type)  # Тип Б
        ws.cell(row=row, column=8, value=birka)  # ID кабеля
        ws.cell(row=row, column=9, value=owner)  # Собственник
        ws.cell(row=row, column=10, value=comment)  # Комментарий
        
        # Save to temporary file and upload to S3
        temp_file = f"/tmp/{journal_filename}"
        wb.save(temp_file)
        with open(temp_file, 'rb') as f:
            s3.upload_fileobj(f, BUCKET_NAME, journal_filename)
        os.remove(temp_file)
        
        # Update reestr in Google Sheets
        last_row_cell = worksheet_reestr.acell('Z1').value
        last_row = int(last_row_cell) if last_row_cell and last_row_cell.isdigit() else 1
        
        now = datetime.now().strftime('%d.%m.%Y %H:%M:%S')
        
        new_row_data = [
            object_number,           # A: OksId
            now,                     # B: Дата
            is_transit,              # C: Транзитность
            object_address,          # D: ТочкаА_адрес
            object_type,             # E: ТочкаА_тип
            object_lat,              # F: ТочкаА_широта
            object_lon,              # G: ТочкаА_долгота
            point_b1,                # H: ТочкаБ_адрес
            point_b_type,            # I: ТочкаБ_тип
            birka,                   # J: ID_кабеля
            owner,                   # K: Владелец
            '',                      # L: ИНН
            '',                      # M: источник
            '',                      # N: адрес
            '',                      # O: телефон
            '',                      # P: email
            comment                  # Q: Комментарий
        ]
        
        worksheet_reestr.insert_row(new_row_data, last_row + 1)
        worksheet_reestr.update_acell('Z1', str(last_row + 1))
        
        # Update completion date in objects sheet
        objects = worksheet_objects.get_all_records()
        for i, obj in enumerate(objects):
            if str(obj.get('№ объекта', '')) == object_number:
                row_num = i + 2  # +2 because header is row 1 and 0-indexed
                worksheet_objects.update_cell(row_num, 4, now)  # Column D
                break
        
        return jsonify({'success': True})
    
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

@app.route('/no_cable', methods=['POST'])
def no_cable():
    try:
        data = request.get_json()
        object_number = data['object_number']
        object_address = data['object_address']
        object_type = data['object_type']
        object_lat = data['object_lat']
        object_lon = data['object_lon']
        
        # Create cable journal with one line
        journal_filename = f"{object_number}_{secure_filename(object_address[:30])}_Кабельный журнал.xlsx"
        
        template_data = worksheet_template.get_all_values()
        wb = load_workbook()
        ws = wb.active
        
        for i in range(min(6, len(template_data))):
            for j in range(len(template_data[i])):
                ws.cell(row=i+1, column=j+1, value=template_data[i][j])
        
        row = 7
        ws.cell(row=row, column=1, value=1)
        ws.cell(row=row, column=2, value=object_address)
        ws.cell(row=row, column=3, value=object_type)
        ws.cell(row=row, column=4, value=object_lat)
        ws.cell(row=row, column=5, value=object_lon)
        ws.cell(row=row, column=6, value='Отсутствует')
        ws.cell(row=row, column=7, value='')
        ws.cell(row=row, column=8, value='Отсутствует')
        ws.cell(row=row, column=9, value='')
        ws.cell(row=row, column=10, value='Кабели не выявлены (отсутствует кабельная инфраструктура)')
        
        temp_file = f"/tmp/{journal_filename}"
        wb.save(temp_file)
        with open(temp_file, 'rb') as f:
            s3.upload_fileobj(f, BUCKET_NAME, journal_filename)
        os.remove(temp_file)
        
        # Update reestr
        last_row_cell = worksheet_reestr.acell('Z1').value
        last_row = int(last_row_cell) if last_row_cell and last_row_cell.isdigit() else 1
        
        now = datetime.now().strftime('%d.%m.%Y %H:%M:%S')
        
        new_row_data = [
            object_number,
            now,
            '',
            object_address,
            object_type,
            object_lat,
            object_lon,
            'Отсутствует',
            '',
            'Отсутствует',
            '',
            '', '', '', '', '',
            'Кабели не выявлены (отсутствует кабельная инфраструктура)'
        ]
        
        worksheet_reestr.insert_row(new_row_data, last_row + 1)
        worksheet_reestr.update_acell('Z1', str(last_row + 1))
        
        # Update completion date
        objects = worksheet_objects.get_all_records()
        for i, obj in enumerate(objects):
            if str(obj.get('№ объекта', '')) == object_number:
                row_num = i + 2
                worksheet_objects.update_cell(row_num, 4, now)
                break
        
        return jsonify({'success': True})
    
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=8080)
